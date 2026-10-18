"""Build ``properties.xlsx`` from the per-site CSV files in ``data/``.

One tab per website with the same columns, an autofilter on every tab and a
summary tab in front.  The workbook itself is not committed to git (see
``storage.py``); CI uploads it to Google Drive and as a workflow artifact.
"""

from __future__ import annotations

import argparse
import logging
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import storage
from sites import SITE_KEYS, load_site
from sites.common import COLUMNS as ALL_COLUMNS, INTERNAL_COLUMNS, NUMERIC_COLUMNS, clean_phone, now_iso

COLUMNS = [column for column in ALL_COLUMNS if column not in INTERNAL_COLUMNS]

LOGGER = logging.getLogger("build-workbook")
OVERVIEW_SHEET = "Overview"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1E915A")
COLUMN_WIDTHS = {
    "Property ID": 12, "Link": 48, "Title": 44, "Deal Type": 18, "Property Type": 18,
    "Status": 18, "Condition": 20, "City": 14, "District": 20, "Neighbourhood": 20,
    "Address": 30, "Area m2": 10, "Yard Area m2": 12, "Rooms": 8, "Bedrooms": 10,
    "Bathrooms": 10, "Floor": 8, "Total Floors": 12, "Price USD": 12, "Price GEL": 12,
    "Price per m2 USD": 14, "Owner Type": 14, "Owner Name": 22, "Phone Number": 20,
    "Agency": 24, "VIP": 12, "Posted": 20, "Updated": 20, "Description": 60,
    "Image": 40, "Timestamp Added": 26, "Detail Status": 12,
}


def write_header(sheet, headers: list[str]) -> None:
    for column_number, header in enumerate(headers, 1):
        cell = sheet.cell(1, column_number, header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def write_site_sheet(workbook: Workbook, title: str, rows: list[dict]) -> None:
    sheet = workbook.create_sheet(title)
    write_header(sheet, COLUMNS)
    for row_number, record in enumerate(rows, 2):
        for column_number, column in enumerate(COLUMNS, 1):
            value = record.get(column, "")
            if column in NUMERIC_COLUMNS and value != "":
                try:
                    value = float(value)
                    value = int(value) if value.is_integer() else value
                except (TypeError, ValueError):
                    pass
            cell = sheet.cell(row_number, column_number, value)
            if column == "Link" and value:
                cell.hyperlink = str(value)
                cell.style = "Hyperlink"
            elif column in {"Property ID", "Phone Number"}:
                cell.number_format = "@"
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 2)}"
    for index, column in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTHS.get(column, 14)


def write_overview(workbook: Workbook, summaries: list[dict]) -> None:
    sheet = workbook.create_sheet(OVERVIEW_SHEET, 0)
    sheet["A1"] = "Real-estate listings monitor"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "One tab per website, all listings unfiltered. Use the column filters in each tab."
    sheet["A3"] = f"Built: {now_iso()}"
    headers = ["Website", "Tab", "Listings", "With phone", "Owner listings", "Newest listing", "Last collected"]
    for column_number, header in enumerate(headers, 1):
        cell = sheet.cell(5, column_number, header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row_number, summary in enumerate(summaries, 6):
        values = [
            summary["site"], summary["sheet"], summary["count"], summary["with_phone"],
            summary["owners"], summary["newest"], summary["collected"],
        ]
        for column_number, value in enumerate(values, 1):
            sheet.cell(row_number, column_number, value)
        link_cell = sheet.cell(row_number, 2)
        link_cell.hyperlink = f"#'{summary['sheet']}'!A1"
        link_cell.style = "Hyperlink"
    sheet.freeze_panes = "A6"
    for index, width in enumerate([16, 16, 12, 12, 16, 22, 26], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def build(output: str | os.PathLike = "properties.xlsx") -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    summaries = []
    for key in SITE_KEYS:
        site = load_site(key)
        rows = list(storage.load_rows(key).values())
        rows.sort(key=lambda row: str(row.get("Timestamp Added") or ""), reverse=True)
        write_site_sheet(workbook, site.SHEET, rows)
        summaries.append({
            "site": site.SHEET,
            "sheet": site.SHEET,
            "count": len(rows),
            "with_phone": sum(1 for row in rows if clean_phone(row.get("Phone Number"))),
            "owners": sum(1 for row in rows if row.get("Owner Type") == "მესაკუთრე"),
            "newest": max((str(row.get("Posted") or row.get("Updated") or "") for row in rows), default=""),
            "collected": max((str(row.get("Timestamp Added") or "") for row in rows), default=""),
        })
        LOGGER.info("Sheet %s: %s rows", site.SHEET, len(rows))
    write_overview(workbook, summaries)

    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_name: str | None = None
    try:
        with tempfile.NamedTemporaryFile(prefix=f".{path.stem}-", suffix=".xlsx", dir=path.parent, delete=False) as handle:
            temp_name = handle.name
        workbook.save(temp_name)
        os.replace(temp_name, path)
        temp_name = None
    finally:
        workbook.close()
        if temp_name and os.path.exists(temp_name):
            os.unlink(temp_name)
    LOGGER.info("Workbook written to %s", path)
    return path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", nargs="?", default="properties.xlsx")
    parser.add_argument("--log-level", default=os.getenv("LOG_LEVEL", "INFO"))
    args = parser.parse_args()
    logging.basicConfig(level=getattr(logging, str(args.log_level).upper(), logging.INFO), format="%(asctime)s %(levelname)s %(message)s")
    build(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
