"""Import the rows collected by the earlier Saburtalo-only monitor.

The previous pipeline stored MyHome rows in a ``Mesakutre`` tab and SS.ge rows
in an ``ss`` tab of ``properties.xlsx`` (still available in git history).
Those owner listings, with the phone numbers it had revealed, are added to
``data/myhome.csv`` and ``data/ssge.csv`` with ``Detail Status = legacy`` so
the newer scope rules never drop them.

    git show <commit>:properties.xlsx > legacy.xlsx
    python import_legacy.py legacy.xlsx
"""

from __future__ import annotations

import argparse

from openpyxl import load_workbook

import storage
from sites.common import OWNER_PHYSICAL, new_row, normalize_id, now_iso

LEGACY_TABS = {"Mesakutre": "myhome", "ss": "ssge"}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook")
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=True)
    for tab, site_key in LEGACY_TABS.items():
        if tab not in workbook.sheetnames:
            print(f"{tab}: tab not found")
            continue
        rows = storage.load_rows(site_key)
        sheet = workbook[tab]
        values = sheet.iter_rows(values_only=True)
        headers = [str(cell) for cell in next(values)]
        added = kept = 0
        for record_values in values:
            record = dict(zip(headers, record_values))
            property_id = normalize_id(record.get("Property ID"))
            if not property_id:
                continue
            if property_id in rows:
                kept += 1
                if not rows[property_id]["Phone Number"] and record.get("Phone Number"):
                    rows[property_id]["Phone Number"] = str(record["Phone Number"])
                continue
            rows[property_id] = new_row(**{
                "Property ID": property_id,
                "Link": record.get("Link") or "",
                "Title": "იყიდება ბინა საბურთალოზე",
                "Deal Type": "იყიდება",
                "Property Type": "ბინა",
                "City": "თბილისი",
                "District": "ვაკე-საბურთალო",
                "Neighbourhood": "საბურთალო",
                "Area m2": record.get("Area") or "",
                "Price USD": record.get("Price") or "",
                "Price per m2 USD": record.get("Price per Sq.m") or "",
                "Owner Type": OWNER_PHYSICAL,
                "Owner Name": record.get("Owner Name") or "",
                "Phone Number": str(record.get("Phone Number") or ""),
                "Posted": str(record.get("Timestamp Added") or "")[:10],
                "Timestamp Added": str(record.get("Timestamp Added") or now_iso()),
                "Detail Status": "legacy",
            })
            added += 1
        storage.save_rows(site_key, rows)
        print(f"{tab} -> data/{site_key}.csv: added {added}, already present {kept}, total {len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
